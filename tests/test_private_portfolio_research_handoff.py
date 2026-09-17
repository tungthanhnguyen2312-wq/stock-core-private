from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

import owner_research_exclusions as ore
import private_portfolio_research_handoff as handoff
from private_portfolio_context import import_workbook


def _workbook(path: Path) -> Path:
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 10, 100])
    trade.append(["2026-01-01", "BBB", "BUY", 5, 50])
    trade.append(["2026-01-02", "BBB", "SELL", 5, 60])       # closed
    trade.append(["2026-01-01", "CCC", "SELL", 3, 100])      # unresolved (no opening buy)
    trade.append(["2026-01-01", "DLQ", "BUY", 7, 20])        # owner-excluded
    workbook.save(path)
    return path


def test_not_available_when_never_imported(tmp_path: Path):
    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    assert artifact["status"] == "NOT_AVAILABLE"
    assert artifact["holdings"] == []


def test_holdings_are_qualified_by_current_position_status(tmp_path: Path):
    import_workbook(workbook_path=_workbook(tmp_path / "wb.xlsx"), portfolio_root=tmp_path)
    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    assert artifact["status"] == "AVAILABLE"
    holdings = {row["ticker"]: row for row in artifact["holdings"]}

    assert holdings["AAA"]["current_position_status"] == "CURRENT_CONFIRMED"
    assert holdings["AAA"]["current_quantity"] == "10"

    assert holdings["BBB"]["current_position_status"] == "CLOSED"
    assert holdings["BBB"]["current_quantity"] == "0"

    assert holdings["CCC"]["current_position_status"] == "CURRENT_POSITION_UNRESOLVED"
    # A reconciliation-unresolved position never carries a quantity, even in the handoff.
    assert holdings["CCC"]["current_quantity"] is None
    assert holdings["CCC"]["current_position_cost_basis_per_share"] is None


def test_owner_excluded_ticker_is_absent_from_holdings_but_named_in_exclusions(tmp_path: Path):
    import_workbook(workbook_path=_workbook(tmp_path / "wb.xlsx"), portfolio_root=tmp_path)
    ore.add_research_exclusion("DLQ", reason="DELISTED_HISTORICAL", excluded_since="2026-09-16", portfolio_root=tmp_path)

    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    tickers = {row["ticker"] for row in artifact["holdings"]}
    assert "DLQ" not in tickers
    assert {row["ticker"] for row in artifact["owner_research_exclusions"]} == {"DLQ"}


def test_console_summary_never_contains_holdings_or_ticker_values(tmp_path: Path):
    import_workbook(workbook_path=_workbook(tmp_path / "wb.xlsx"), portfolio_root=tmp_path)
    ore.add_research_exclusion("DLQ", reason="DELISTED_HISTORICAL", portfolio_root=tmp_path)
    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    summary = handoff.public_console_summary(artifact, destination=tmp_path / "out.json")

    assert "holdings" not in summary
    assert "owner_research_exclusions" not in summary
    assert summary["holding_count"] == 3  # AAA, BBB, CCC -- DLQ excluded
    assert summary["owner_research_exclusion_count"] == 1
    for value in summary.values():
        assert "AAA" not in str(value) and "BBB" not in str(value) and "CCC" not in str(value) and "DLQ" not in str(value)


def test_written_artifact_never_lands_inside_the_repository(tmp_path: Path):
    import_workbook(workbook_path=_workbook(tmp_path / "wb.xlsx"), portfolio_root=tmp_path)
    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    destination = handoff.write_private_artifact(artifact, portfolio_root=tmp_path)
    assert destination.is_file()
    assert tmp_path in destination.parents


def _multi_account_workbook(path: Path) -> Path:
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 10, 100])
    account = workbook.create_sheet("AccountSnapshot")
    account.append(["as_of", "account_alias", "broker", "cash_investable"])
    account.append(["2026-02-01", "105C793100", "SSI", 1000])
    account.append(["2026-02-01", "001302384", "VNDIRECT", 2000])
    workbook.save(path)
    return path


def test_investment_accounts_context_never_leaks_a_raw_account_alias(tmp_path: Path):
    """Section 10: the real workbook's own account_alias can literally be a brokerage account
    number (e.g. "105C793100") -- this externally-uploadable artifact must never carry it."""
    import_workbook(workbook_path=_multi_account_workbook(tmp_path / "wb.xlsx"), portfolio_root=tmp_path)
    artifact = handoff.build_artifact(portfolio_root=tmp_path)
    context = artifact["investment_accounts_context"]

    assert context["status"] == "AGGREGATED"
    assert context["account_count"] == 2
    labels = {account["handoff_account_label"] for account in context["accounts"]}
    assert labels == {"ACCOUNT_1", "ACCOUNT_2"}
    brokers = {account["broker"] for account in context["accounts"]}
    assert brokers == {"SSI", "VNDIRECT"}
    assert "105C793100" not in str(artifact)
    assert "001302384" not in str(artifact)
