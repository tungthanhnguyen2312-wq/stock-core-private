from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

import private_portfolio_context as portfolio_context
from private_portfolio_context import (
    CURRENT_COST_BASIS_METHOD,
    CURRENT_POSITION_STATUS_CLOSED,
    CURRENT_POSITION_STATUS_CONFIRMED,
    CURRENT_POSITION_STATUS_UNRESOLVED,
    LIFETIME_BREAKEVEN_METHOD,
    PortfolioImportError,
    import_workbook,
    portfolio_status,
    public_import_summary,
    public_status_summary,
)
from stocklookup import main


def _workbook(path: Path, *, inconsistent: bool = False, reopened: bool = False, account_policy: bool = False) -> Path:
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price", "Fee", "Tax"])
    if reopened:
        trade.append(["2026-01-01", "AAA", "BUY", 1, 10, 0, 0])
        trade.append(["2026-01-02", "AAA", "SELL", 1, 11, 0, 0])
        trade.append(["2026-01-03", "AAA", "BUY", 2, 12, 0, 0])
    else:
        trade.append(["2026-01-01", "AAA", "BUY", 10, 100, 1, 0])
        trade.append(["2026-01-02", "AAA", "SELL", 4, 120, 1, 2])
        trade.append(["2026-01-03", "AAA", "BUY", 2, 90, 0, 0])
        if inconsistent:
            trade.append(["2026-01-04", "AAA", "SELL", 20, 100, 0, 0])

    dividend = workbook.create_sheet("Dividend")
    dividend.append(["Date", "Ticker", "Type", "Cash Amount", "Stock Quantity"])
    if not reopened:
        dividend.append(["2026-01-04", "AAA", "cash dividend", 20, None])
        dividend.append(["2026-01-05", "AAA", "stock dividend", None, 2])

    money = workbook.create_sheet("Money")
    money.append(["Date", "Ticker", "Type", "Amount"])
    if not reopened:
        money.append(["2026-01-06", "AAA", "tax", 5])

    margin = workbook.create_sheet("margin")
    margin.append(["Date", "Ticker", "Type", "Amount"])
    if not reopened:
        margin.append(["2026-01-06", "AAA", "interest", 3])
        margin.append(["2026-01-06", None, "interest", 7])

    total = workbook.create_sheet("Total")
    total.append(["Ticker", "Quantity"])
    total.append(["AAA", 1 if inconsistent else 10 if not reopened else 2])

    if account_policy:
        account = workbook.create_sheet("AccountSnapshot")
        account.append(["As Of Date", "Currency", "Cash Available", "Margin Debt", "Net Asset Value"])
        account.append(["2026-01-06", "VND", 1000, 200, 5000])
        policy = workbook.create_sheet("PortfolioPolicy")
        policy.append(["As Of Date", "Currency", "Max Single Position Weight", "Max Margin Debt To NAV"])
        policy.append(["2026-01-06", "VND", 0.25, 0.2])

    workbook.save(path)
    return path


def _position(result: dict, ticker: str = "AAA") -> dict:
    return next(row for row in result["snapshot"]["positions"] if row["ticker"] == ticker)


def test_trade_dividend_margin_and_cost_semantics_are_deterministic(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "synthetic.xlsx"), portfolio_root=tmp_path / "private")
    position = _position(result)

    assert position["current_quantity"] == "10"
    assert position["current_position_status"] == CURRENT_POSITION_STATUS_CONFIRMED
    assert position["position_episode_holding_days"] == 5
    assert position["current_position_carrying_cost"] == "780.6"
    assert position["current_position_cost_basis_per_share"] == "78.06"
    assert position["current_position_cost_basis_method"] == CURRENT_COST_BASIS_METHOD
    assert position["purchase_count"] == 2
    assert position["sale_count"] == 1
    assert position["cash_dividends"] == "20"
    assert position["stock_distribution_quantity"] == "2"
    assert position["ticker_financing_cost"] == "3"
    assert position["ticker_fee_or_tax_cashflow"] == "5"
    assert position["lifetime_net_cash_outflow"] == "692"
    assert position["lifetime_cash_recovery_breakeven"] == "69.2"
    assert position["lifetime_cash_recovery_breakeven_method"] == LIFETIME_BREAKEVEN_METHOD
    assert position["realized_pnl"] == "76.6"
    assert position["unrealized_pnl"] is None
    assert result["snapshot"]["account_level_unallocated_margin_cost"] == "7"


def test_reopened_position_has_new_episode_without_rewriting_lifetime_history(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "reopened.xlsx", reopened=True), portfolio_root=tmp_path / "private")
    position = _position(result)

    assert position["current_quantity"] == "2"
    assert position["position_episode_start_date"] == "2026-01-03"
    assert position["position_episode_holding_days"] == 0
    assert position["first_acquisition_date"] == "2026-01-01"
    assert position["last_acquisition_date"] == "2026-01-03"
    assert position["purchase_count"] == 2
    assert position["sale_count"] == 1


def test_oversize_sell_never_leaves_a_phantom_current_holding(tmp_path: Path):
    """PERSONAL_DECISION_INPUT_TRUTH_V1: a SELL that exceeds the derived running holding must
    never leave a stale positive `current_quantity` readable as a confirmed current holding --
    the reconstructed quantity is withheld entirely and the position is flagged
    CURRENT_POSITION_UNRESOLVED, distinct from both a genuine current holding and a closed one."""
    result = import_workbook(workbook_path=_workbook(tmp_path / "inconsistent.xlsx", inconsistent=True), portfolio_root=tmp_path / "private")
    position = _position(result)

    assert position["current_position_status"] == CURRENT_POSITION_STATUS_UNRESOLVED
    assert position["current_quantity"] is None
    assert position["realized_pnl"] is None
    assert position["realized_pnl_status"] == "UNRESOLVED_RECONCILIATION_WARNING"
    codes = result["snapshot"]["reconciliation"]["warning_counts"]
    assert codes["SELL_QUANTITY_EXCEEDS_DERIVED_HOLDING"] == 1
    assert codes["TOTAL_VIEW_HINT_QUANTITY_MISMATCH"] == 1
    # HISTORICAL_EVENT_LEDGER_STATE is untouched: every real event, including the oversize
    # SELL itself, is still retained verbatim in the ledger despite the position being unresolved.
    sell_events = [event for event in result["ledger"]["events"] if event["event_type"] == "SELL" and event["ticker"] == "AAA"]
    assert len(sell_events) == 2
    assert result["snapshot"]["current_position_status_counts"] == {CURRENT_POSITION_STATUS_UNRESOLVED: 1}
    # A reconciliation-blocked position must never count as an actual current position.
    assert public_import_summary(result)["current_position_count"] == 0


def test_missing_historical_opening_position_sell_with_no_prior_buy_is_unresolved(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    # No opening BUY at all for this ticker -- the very first event is a SELL.
    trade.append(["2026-01-01", "ZZZ", "SELL", 5, 100])
    path = tmp_path / "no-opening-position.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    position = _position(result, ticker="ZZZ")

    assert position["current_position_status"] == CURRENT_POSITION_STATUS_UNRESOLVED
    assert position["current_quantity"] is None
    assert "SELL_QUANTITY_EXCEEDS_DERIVED_HOLDING" in result["snapshot"]["reconciliation"]["warning_counts"]


def test_fully_sold_down_position_is_closed_not_current(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "reopened.xlsx", reopened=True), portfolio_root=tmp_path / "private")
    # `_workbook(reopened=True)` ends with a positive quantity (BUY 1, SELL 1, BUY 2); build a
    # fresh fixture that instead sells the full position back down to zero.
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 10, 100])
    trade.append(["2026-01-02", "AAA", "SELL", 10, 110])
    path = tmp_path / "closed-position.xlsx"
    workbook.save(path)

    closed_result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private-closed")
    position = _position(closed_result)

    assert position["current_position_status"] == CURRENT_POSITION_STATUS_CLOSED
    assert position["current_quantity"] == "0"
    assert closed_result["snapshot"]["reconciliation"]["status"] == "RECONCILED"
    assert public_import_summary(closed_result)["current_position_count"] == 0
    assert public_status_summary(portfolio_status(portfolio_root=tmp_path / "private-closed"))["current_position_count"] == 0
    # Sanity: the CURRENT_CONFIRMED position from the unrelated reopened-position fixture above
    # is unaffected by any of this (confirms the fix is per-ticker, not a global state leak).
    assert _position(result)["current_position_status"] == CURRENT_POSITION_STATUS_CONFIRMED


def test_current_position_count_excludes_closed_and_unresolved_positions(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 10, 100])  # stays CURRENT_CONFIRMED
    trade.append(["2026-01-01", "BBB", "BUY", 5, 100])
    trade.append(["2026-01-02", "BBB", "SELL", 5, 110])  # fully closed -> CLOSED
    trade.append(["2026-01-01", "CCC", "SELL", 3, 100])  # no opening buy -> UNRESOLVED
    path = tmp_path / "mixed-positions.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    summary = public_import_summary(result)

    assert summary["current_position_count"] == 1
    assert summary["current_position_status_counts"] == {
        CURRENT_POSITION_STATUS_CLOSED: 1,
        CURRENT_POSITION_STATUS_CONFIRMED: 1,
        CURRENT_POSITION_STATUS_UNRESOLVED: 1,
    }


def test_absent_account_and_policy_are_explicit_not_zero_filled(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "absent.xlsx"), portfolio_root=tmp_path / "private")

    assert result["account_snapshot"]["status"] == "NOT_PROVIDED"
    assert set(result["account_snapshot"]["fields"].values()) == {None}
    assert result["policy"]["status"] == "SYSTEM_DEFAULTS_APPLIED"
    assert result["policy"]["owner_policy_status"] == "NOT_PROVIDED"
    assert result["policy"]["owner_supplied_fields"]["max_single_position_weight"] is None
    assert result["policy"]["effective_fields"]["max_single_position_weight"] == "0.30"
    assert result["policy"]["field_provenance"]["max_single_position_weight"]["effective_source"] == "SYSTEM_DEFAULT_POLICY_V2"
    assert result["system_default_policy"]["policy_version"] == "SYSTEM_DEFAULT_POLICY_V2"
    assert result["system_default_policy"]["default_fields"] == {
        "risk_budget_per_investment_decision_to_nav": "0.01",
        "max_single_position_weight": "0.30",
        "max_sector_weight": "0.45",
        "max_gross_exposure_to_nav": "1.15",
        "max_margin_debt_to_nav": "0.15",
        "minimum_cash_reserve_to_nav": "0.05",
        "max_margin_rate_percent_for_new_leveraged_exposure": "15.0",
        # PORTFOLIO_AWARE_DECISION_AND_RISK_SIZING_V1 additions (2026-09-09).
        "probe_risk_budget_multiplier": "0.50",
        "tactical_margin_holding_days": "30",
        "min_net_reward_risk_for_margin": "2.0",
        "strong_net_reward_risk_for_max_margin": "3.0",
        "max_financing_cost_fraction_of_gross_upside": "0.20",
    }


def test_account_snapshot_and_policy_are_optional_private_contracts(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "policy.xlsx", account_policy=True), portfolio_root=tmp_path / "private")

    assert result["account_snapshot"]["status"] == "PROVIDED"
    assert result["account_snapshot"]["fields"]["cash_available"] == "1000"
    assert result["policy"]["status"] == "MIXED_OWNER_AND_SYSTEM_DEFAULTS"
    assert result["policy"]["fields"]["max_single_position_weight"] == "0.25"
    assert result["policy"]["field_provenance"]["max_single_position_weight"]["effective_source"] == "OWNER_WORKBOOK"
    assert result["policy"]["fields"]["minimum_cash_reserve_to_nav"] == "0.05"


def test_account_snapshot_retains_distinct_margin_availability_and_rate_without_debt_inference(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 100])
    account = workbook.create_sheet("AccountSnapshot")
    account.append(["As Of Date", "Cash Available", "Net Asset Value", "Margin Available Minimum", "Margin Available Maximum", "Annual Margin Rate Percent"])
    account.append(["2026-01-02", 5000000, 10000000, 1000000, 7000000, 13.5])
    # The workbook leaves all monetary cells unformatted (the Excel default).
    assert account["B2"].number_format == "General"
    path = tmp_path / "margin-account.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    fields = result["account_snapshot"]["fields"]

    assert fields["cash_available"] == "5000000"
    assert fields["margin_available_minimum"] == "1000000"
    assert fields["margin_available_maximum"] == "7000000"
    assert fields["annual_margin_rate_percent"] == "13.5"
    assert fields["margin_debt"] is None
    assert result["account_snapshot"]["field_semantics"]["annual_margin_rate_percent"]["unit_or_format"] == "PERCENT_PER_ANNUM_NUMBER"
    assert result["account_snapshot"]["field_semantics"]["margin_debt"]["unit_or_format"] == "VND"


def test_legacy_margin_sheet_fields_are_reused_without_duplicate_account_snapshot_input(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 100])
    margin = workbook.create_sheet("margin")
    margin.append(["Field", "Value"])
    margin.append(["Margin Available Minimum", 1000000])
    margin.append(["Margin Available Maximum", 7000000])
    margin.append(["Annual Margin Rate Percent", 13.5])
    path = tmp_path / "legacy-margin-fields.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    fields = result["account_snapshot"]["fields"]

    assert result["account_snapshot"]["source"] == {"sheets": ["margin"]}
    assert fields["margin_available_minimum"] == "1000000"
    assert fields["margin_available_maximum"] == "7000000"
    assert fields["annual_margin_rate_percent"] == "13.5"
    assert fields["margin_debt"] is None
    assert result["account_snapshot"]["field_provenance"]["margin_available_minimum"] == {"selected_source": "margin", "value_origin": "OWNER_WORKBOOK_FIELD"}
    assert result["account_snapshot"]["authority_boundary"]["margin_availability_is_not_margin_debt"] is True


def test_money_amount_is_vnd_by_field_identity_without_format_or_magnitude_inference(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 100])
    money = workbook.create_sheet("Money")
    money.append(["Date", "Type", "Amount"])
    money.append(["2026-01-02", "cash transfer", 1234567])
    assert money["C2"].number_format == "General"
    path = tmp_path / "unformatted-vnd.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    movement = next(event for event in result["ledger"]["events"] if event["event_type"] == "MONEY_MOVEMENT")

    assert movement["gross_amount"] == "1234567"
    assert movement["monetary_currency"] == "VND"
    assert movement["monetary_unit_basis"] == "WORKBOOK_FIELD_IDENTITY_VND_NO_FORMAT_OR_MAGNITUDE_INFERENCE"
    assert result["ledger"]["authority_boundary"]["money_and_margin_currency_inference_from_cell_format_or_magnitude"] == "PROHIBITED"


def test_partial_owner_policy_override_keeps_source_effective_values_and_provenance(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 100])
    policy = workbook.create_sheet("PortfolioPolicy")
    policy.append(["Max Single Position Weight"])
    policy.append([0.22])
    path = tmp_path / "partial-policy.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    policy_result = result["policy"]

    assert policy_result["status"] == "MIXED_OWNER_AND_SYSTEM_DEFAULTS"
    assert policy_result["owner_supplied_fields"]["max_single_position_weight"] == "0.22"
    assert policy_result["effective_fields"]["max_single_position_weight"] == "0.22"
    assert policy_result["field_provenance"]["max_single_position_weight"]["effective_source"] == "OWNER_WORKBOOK"
    assert policy_result["owner_supplied_fields"]["max_margin_debt_to_nav"] is None
    assert policy_result["effective_fields"]["max_margin_debt_to_nav"] == "0.15"
    assert policy_result["field_provenance"]["max_margin_debt_to_nav"]["effective_source"] == "SYSTEM_DEFAULT_POLICY_V2"
    assert policy_result["effective_fields"]["max_margin_rate_percent_for_new_leveraged_exposure"] == "15.0"
    assert policy_result["authority_boundary"]["existing_positions_above_policy_caps_are_not_automatic_sell_instructions"] is True


def test_reimport_is_content_addressed_and_status_does_not_reopen_workbook(tmp_path: Path):
    workbook = _workbook(tmp_path / "idempotent.xlsx")
    root = tmp_path / "private"
    first = import_workbook(workbook_path=workbook, portfolio_root=root)
    second = import_workbook(workbook_path=workbook, portfolio_root=root)
    status = portfolio_status(portfolio_root=root)

    assert first["status"] == "IMPORTED"
    assert second["status"] == "REUSED_IDENTICAL"
    assert first["manifest"]["artifact_identity"] == second["manifest"]["artifact_identity"]
    assert status["status"] == "READY"
    assert status["manifest"]["artifact_identity"] == first["manifest"]["artifact_identity"]
    assert json.loads((root / "latest_import.json").read_text(encoding="utf-8"))["workbook_sha256"] == first["manifest"]["workbook_sha256"]


def test_corrective_import_layout_preserves_a_prior_same_workbook_sha_artifact(tmp_path: Path):
    workbook = _workbook(tmp_path / "prior-layout.xlsx")
    root = tmp_path / "private"
    legacy_directory = root / "imports" / portfolio_context._sha256_file(workbook)
    legacy_directory.mkdir(parents=True)
    legacy_artifact = legacy_directory / "portfolio_event_ledger_v1.json"
    legacy_artifact.write_text('{"legacy":"immutable"}\n', encoding="utf-8")

    result = import_workbook(workbook_path=workbook, portfolio_root=root)

    assert result["status"] == "IMPORTED"
    assert result["private_artifact_directory"] != legacy_directory
    assert legacy_artifact.read_text(encoding="utf-8") == '{"legacy":"immutable"}\n'
    assert result["manifest"]["import_layout_version"] == "PORTFOLIO_CONTEXT_IMPORT_LAYOUT_V6"
    assert result["private_artifact_directory"].parent == legacy_directory


def test_owner_cli_portfolio_branch_uses_private_summary_without_daily_preflight(tmp_path: Path, capsys):
    workbook = _workbook(tmp_path / "cli.xlsx")
    root = tmp_path / "private"

    assert main(["portfolio", "import", "--workbook", str(workbook), "--portfolio-root", str(root)]) == 0
    imported = json.loads(capsys.readouterr().out)
    assert imported["status"] == "IMPORTED"
    assert imported["provider_calls"] == "NOT_USED"
    assert imported["daily_run"] == "NOT_USED"
    assert "positions" not in imported
    assert "events" not in imported

    assert main(["portfolio", "status", "--portfolio-root", str(root)]) == 0
    status = json.loads(capsys.readouterr().out)
    assert status["status"] == "READY"
    assert status["import_manifest_identity"] == imported["import_manifest_identity"]


def test_workbook_inside_repository_boundary_is_refused_even_when_synthetic(tmp_path: Path, monkeypatch):
    synthetic_repository = tmp_path / "repository"
    synthetic_repository.mkdir()
    workbook = _workbook(synthetic_repository / "synthetic.xlsx")
    monkeypatch.setattr(portfolio_context, "REPOSITORY_ROOT", synthetic_repository)

    with pytest.raises(PortfolioImportError, match="MUST_LIVE_OUTSIDE_REPOSITORY"):
        import_workbook(workbook_path=workbook, portfolio_root=tmp_path / "private")


def test_header_row_below_a_title_row_is_detected_and_labels_are_normalized(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["TRADE HISTORY (DRAFT EXPORT)"])
    trade.append([])
    trade.append([" Date ", " TICKER ", "Side:", "QTY.", "Price"])
    trade.append(["2026-01-01", "aaa", "buy", 5, 100])
    path = tmp_path / "non-row1-header.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")

    assert result["ledger"]["reconciliation_warnings"] == []
    assert _position(result)["current_quantity"] == "5"


def test_vietnamese_d_with_stroke_folds_to_plain_d_not_dropped(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Ngày giao dịch", "Mã CK", "Hành động", "KL khớp", "Giá khớp"])
    trade.append(["2026-01-01", "AAA", "Mua", 5, 100])
    path = tmp_path / "d-stroke.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")

    assert result["ledger"]["reconciliation_warnings"] == []
    assert _position(result)["current_quantity"] == "5"


def test_dividend_row_falls_back_to_record_or_exrights_date_when_execution_date_is_blank(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 100])
    dividend = workbook.create_sheet("Dividend")
    dividend.append(["Cổ phiếu", "Tiền cổ tức", "Ngày GDKHQ", "Ngày ĐKCC", "Ngày thực hiện"])
    # Not yet settled: only the ex-rights and record dates are filled in.
    dividend.append(["AAA", 500, "2026-02-01", "2026-02-03", None])
    path = tmp_path / "date-fallback.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    codes = {warning["code"] for warning in result["ledger"]["reconciliation_warnings"]}
    cash_dividend = next(event for event in result["ledger"]["events"] if event["event_type"] == "CASH_DIVIDEND")

    assert "DIVIDEND_ROW_INVALID" not in codes
    # Record date (Ngay DKCC) is preferred over ex-rights date when the
    # execution date is absent -- it is the later, more settled milestone.
    assert cash_dividend["effective_date"] == "2026-02-03"


def test_real_style_vietnamese_trade_dividend_money_and_margin_layouts_are_recognized(tmp_path: Path):
    """Synthetic fixture reproducing the real workbook's structure (sheet
    names, header-row shape, exact column labels) with fabricated data only
    -- never a real holding, quantity, price, or amount."""
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append([
        "STT", "Ngày giao dịch", "Ngày giao dịch2", "Ngày thanh toán", "Thời gian khớp",
        "Mua/Bán", "Loại lệnh", "Mã CK", "KL đặt", "Giá đặt", "KL khớp", "Giá khớp",
        "Giá trị khớp", "Phí GD", "Thuế TNCN", "Thành tiền", "Số tài khoản", "Subject",
        "Khóa chống trùng", "Tháng", "Năm",
    ])
    trade.append([
        1, "2026-01-01", "2026-01-01", "2026-01-03", None, "Mua", "LO", "AAA",
        10, 100, 10, 100, 1000, 2, 0, 1002, None, None, None, 1, 2026,
    ])

    dividend = workbook.create_sheet("Dividend")
    dividend.append([
        "STT", "Cổ phiếu", "Loại Sự kiện", "Số tiền thực nhận\n trừ thuế",
        "Số CK\n được nhận/được mua", "Số CK\nhưởng quyền", "Tỉ lệ", "Tiền cổ tức",
        "Số tiền được nhận", "Số CK đã \nđăng ký mua", "Giá mua", "Số tiền đã nộp",
        "Tài khoản", "Ngày GDKHQ", "Ngày ĐKCC", "Ngày thực hiện", "Nội dung sự kiện",
        "Ảnh hưởng KL", "Dòng tiền quyền/cổ tức",
    ])
    # Matches the real workbook's own pattern: "So tien thuc nhan tru thue"
    # (col 4) reliably carries the cash-dividend amount; "Tien co tuc" (col 8)
    # is typically blank on real rows.
    dividend.append([2, "AAA", "Cổ tức tiền mặt", 500, None, None, None, None, None, None, None, None, None, "2026-02-01", "2026-02-01", "2026-02-01", None, None, None])
    dividend.append([3, "AAA", "Cổ phiếu thưởng", None, 3, 3.7, None, None, None, None, None, None, None, "2026-02-02", "2026-02-02", "2026-02-02", None, None, None])

    money = workbook.create_sheet("Money")
    money.append(["STT", "Ngày giao dịch", "Hành động", "Công ty CK", "Tài khoản", "Số tiền", "Remark"])
    money.append([4, "2026-01-10", "Rút tiền", None, None, 200, None])

    margin = workbook.create_sheet("margin")
    margin.append([])
    margin.append([None, "2026-01-05", 1000000, "vay hsbc nạp vào"])
    margin.append([None, "Tổng nạp", 1000000])

    path = tmp_path / "real-style-layout.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    ledger = result["ledger"]
    codes = {warning["code"] for warning in ledger["reconciliation_warnings"]}

    assert "TRADE_HEADER_UNRECOGNIZED" not in codes
    assert "DIVIDEND_HEADER_UNRECOGNIZED" not in codes
    assert "MONEY_HEADER_UNRECOGNIZED" not in codes
    assert "MARGIN_HEADER_UNRECOGNIZED" not in codes
    assert "MARGIN_LEGACY_POSITIONAL_LAYOUT_USED" in codes
    assert len(ledger["events"]) > 0
    assert len(result["snapshot"]["positions"]) > 0

    position = _position(result)
    assert position["cash_dividends"] == "500"
    assert position["stock_distribution_quantity"] == "3"
    buy_event = next(event for event in ledger["events"] if event["event_type"] == "BUY")
    assert buy_event["gross_amount"] == "1000"
    money_movements = [event for event in ledger["events"] if event["event_type"] == "MONEY_MOVEMENT"]
    assert {event["gross_amount"] for event in money_movements} == {"200", "1000000"}


def test_fractional_corporate_action_entitlement_is_not_a_reconciliation_mismatch(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    trade.append(["2026-01-01", "AAA", "BUY", 100, 10])
    dividend = workbook.create_sheet("Dividend")
    dividend.append(["Cổ phiếu", "Loại Sự kiện", "Số CK\n được nhận/được mua", "Số CK\nhưởng quyền", "Ngày thực hiện"])
    # The theoretical entitlement (13.7) is fractional; the broker-settled
    # quantity actually added to the position (13) is a whole number.
    dividend.append(["AAA", "Cổ phiếu thưởng", 13, 13.7, "2026-02-01"])
    path = tmp_path / "fractional-entitlement.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    position = _position(result)
    codes = {warning["code"] for warning in result["ledger"]["reconciliation_warnings"]}

    assert position["current_quantity"] == "113"
    assert position["stock_distribution_quantity"] == "13"
    assert not any("ENTITLEMENT" in code or "FRACTIONAL" in code for code in codes)


def test_zero_events_with_all_source_headers_recognized_fails_closed_not_ready(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    dividend = workbook.create_sheet("Dividend")
    dividend.append(["Date", "Ticker", "Cash Amount"])
    money = workbook.create_sheet("Money")
    money.append(["Date", "Amount"])
    margin = workbook.create_sheet("margin")
    margin.append(["Date", "Amount"])
    path = tmp_path / "zero-activity.xlsx"
    workbook.save(path)

    root = tmp_path / "private"
    result = import_workbook(workbook_path=path, portfolio_root=root)
    codes = {warning["code"] for warning in result["ledger"]["reconciliation_warnings"]}
    assert "ZERO_EVENTS_DESPITE_RECOGNIZED_SOURCES" in codes
    assert result["ledger"]["events"] == []

    status = portfolio_status(portfolio_root=root)
    assert status["status"] == "RECONCILIATION_INCOMPLETE"
    assert "ZERO_EVENTS_DESPITE_RECOGNIZED_SOURCES" in status["reason_codes"]


def test_labeled_net_trade_amount_does_not_double_count_fee_or_tax(tmp_path: Path):
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Net Amount", "Fee", "Tax"])
    trade.append(["2026-01-01", "AAA", "BUY", 1, 105, 3, 2])
    trade.append(["2026-01-02", "AAA", "SELL", 1, 115, 3, 2])
    path = tmp_path / "net-amount.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    events = result["ledger"]["events"]
    position = _position(result)
    assert [event["gross_amount"] for event in events] == ["100", "120"]
    assert position["realized_pnl"] == "10"


# ── PRIVATE_MULTI_BROKER_INVESTMENT_ACCOUNT_CONTEXT_V1 ───────────────────────────────────────


def test_legacy_single_unaliased_account_row_imports_as_one_legacy_account(tmp_path: Path):
    """Section 6: an existing legacy workbook (one AccountSnapshot row, no account_alias column
    populated) must still import successfully, as exactly one account with an explicit
    legacy/unspecified identity -- never zero accounts, never a fabricated alias."""
    result = import_workbook(workbook_path=_workbook(tmp_path / "legacy.xlsx", account_policy=True), portfolio_root=tmp_path / "private")
    accounts = result["investment_accounts"]["accounts"]

    assert len(accounts) == 1
    assert accounts[0]["account_id"] == portfolio_context.LEGACY_UNSPECIFIED_ACCOUNT_ID
    assert accounts[0]["account_id_basis"] == "LEGACY_UNSPECIFIED_SINGLE_ACCOUNT"
    assert accounts[0]["fields"]["cash_available"] == "1000"
    assert accounts[0]["fields"]["margin_debt"] == "200"
    assert accounts[0]["fields"]["broker_reported_nav"] == "5000"
    aggregate = result["investment_accounts_portfolio_context"]
    assert aggregate["status"] == "AGGREGATED"
    assert aggregate["account_count"] == 1
    assert aggregate["totals"]["total_broker_cash"] == "1000"
    assert aggregate["totals"]["total_margin_debt"] == "200"
    assert aggregate["totals"]["total_broker_reported_nav"] == "5000"
    # The pre-existing single-account contract is completely unaffected by the new one.
    assert result["account_snapshot"]["fields"]["cash_available"] == "1000"


def _multi_account_workbook(path: Path) -> Path:
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price", "Account"])
    trade.append(["2026-01-01", "AAA", "BUY", 10, 100, "ACC-A"])
    trade.append(["2026-01-02", "AAA", "BUY", 5, 100, "ACC-B"])
    trade.append(["2026-01-03", "BBB", "BUY", 4, 50, None])  # no per-row account -> unresolved

    account = workbook.create_sheet("AccountSnapshot")
    account.append(["as_of", "account_alias", "broker", "cash_investable", "margin_debt", "broker_nav"])
    account.append(["2026-02-01", "ACC-A", "SSI", 1000, 100, 5000])
    account.append(["2026-02-01", "ACC-B", "VNDIRECT", 2000, 0, 8000])
    workbook.save(path)
    return path


def test_multiple_aliased_accounts_and_consistent_as_of_aggregate_deterministically(tmp_path: Path):
    """Section 3/4: two distinct broker accounts, each independently represented, aggregate
    deterministically when their as-of dates and currency agree."""
    result = import_workbook(workbook_path=_multi_account_workbook(tmp_path / "multi.xlsx"), portfolio_root=tmp_path / "private")
    accounts = {account["account_id"]: account for account in result["investment_accounts"]["accounts"]}

    assert set(accounts) == {"ACC-A", "ACC-B"}
    assert accounts["ACC-A"]["fields"]["broker"] == "SSI"
    assert accounts["ACC-A"]["fields"]["cash_available"] == "1000"
    assert accounts["ACC-B"]["fields"]["broker"] == "VNDIRECT"
    assert accounts["ACC-B"]["fields"]["cash_available"] == "2000"

    aggregate = result["investment_accounts_portfolio_context"]
    assert aggregate["status"] == "AGGREGATED"
    assert aggregate["account_count"] == 2
    assert aggregate["as_of_consistency"] == "CONSISTENT"
    # Each account's own cash is summed exactly once -- no double counting.
    assert aggregate["totals"]["total_broker_cash"] == "3000"
    assert aggregate["totals"]["total_margin_debt"] == "100"
    assert aggregate["totals"]["total_broker_reported_nav"] == "13000"

    # Global CURRENT_POSITION_TRUTH is untouched by per-account attribution: AAA's total quantity
    # (10 + 5 across both accounts) is still the single authoritative current_quantity/status.
    aaa_position = _position(result, "AAA")
    assert aaa_position["current_quantity"] == "15"
    assert aaa_position["current_position_status"] == CURRENT_POSITION_STATUS_CONFIRMED

    # Section 5: the same ticker (AAA) held in two accounts keeps two separate lineage rows.
    aaa_attribution = {row["account_id"]: row for row in _position(result, "AAA")["account_attribution"]}
    assert aaa_attribution["ACC-A"]["current_quantity"] == "10"
    assert aaa_attribution["ACC-B"]["current_quantity"] == "5"
    assert aaa_attribution["ACC-A"]["attribution_status"] == "ATTRIBUTED"
    # BBB's only event carries no per-row account -> unresolved, never guessed.
    bbb_attribution = _position(result, "BBB")["account_attribution"]
    assert len(bbb_attribution) == 1
    assert bbb_attribution[0]["account_id"] == portfolio_context.ACCOUNT_ATTRIBUTION_UNRESOLVED
    assert bbb_attribution[0]["attribution_status"] == "ACCOUNT_ATTRIBUTION_UNRESOLVED"

    summary = result["snapshot"]["position_account_attribution_summary"]
    assert summary["tickers_with_attributed_account"] == 1
    assert summary["tickers_unresolved_attribution_only"] == 1


def test_mismatched_account_as_of_dates_fail_closed_never_silently_summed(tmp_path: Path):
    """Section 4: accounts with different as-of dates must never be silently summed together."""
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    account = workbook.create_sheet("AccountSnapshot")
    account.append(["as_of", "account_alias", "cash_investable"])
    account.append(["2026-02-01", "ACC-A", 1000])
    account.append(["2026-01-15", "ACC-B", 2000])
    path = tmp_path / "mismatched-as-of.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    aggregate = result["investment_accounts_portfolio_context"]

    assert aggregate["status"] == "PARTIAL_MULTI_ACCOUNT_MISMATCH"
    assert aggregate["as_of_consistency"] == "MULTI_ACCOUNT_AS_OF_MISMATCH"
    assert "MULTI_ACCOUNT_AS_OF_MISMATCH" in aggregate["reason_codes"]
    assert aggregate["totals"]["total_broker_cash"] is None
    assert aggregate["totals"]["total_broker_reported_nav"] is None


def test_incomplete_field_coverage_never_manufactures_a_partial_total(tmp_path: Path):
    """Section 4: a total is UNKNOWN, never a partial/manufactured figure, when any qualified
    account omits that specific field (e.g. only one account reports broker NAV)."""
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    account = workbook.create_sheet("AccountSnapshot")
    account.append(["as_of", "account_alias", "cash_investable", "broker_nav"])
    account.append(["2026-02-01", "ACC-A", 1000, 5000])
    account.append(["2026-02-01", "ACC-B", 2000, None])
    path = tmp_path / "incomplete-coverage.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    aggregate = result["investment_accounts_portfolio_context"]

    assert aggregate["status"] == "AGGREGATED"
    assert aggregate["totals"]["total_broker_cash"] == "3000"
    assert aggregate["totals"]["total_broker_reported_nav"] is None
    assert aggregate["total_field_basis"]["total_broker_reported_nav"]["reason"] == "INCOMPLETE_ACCOUNT_COVERAGE_FOR_FIELD"


def test_ambiguous_unaliased_multi_row_never_collapses_or_guesses_identity(tmp_path: Path):
    """Two AccountSnapshot rows both missing an alias is genuinely ambiguous -- never silently
    treated as one legacy account, never silently dropped."""
    workbook = Workbook()
    trade = workbook.active
    trade.title = "Trade"
    trade.append(["Date", "Ticker", "Side", "Quantity", "Price"])
    account = workbook.create_sheet("AccountSnapshot")
    account.append(["as_of", "account_alias", "cash_investable"])
    account.append(["2026-02-01", None, 1000])
    account.append(["2026-02-01", None, 2000])
    path = tmp_path / "ambiguous.xlsx"
    workbook.save(path)

    result = import_workbook(workbook_path=path, portfolio_root=tmp_path / "private")
    accounts = result["investment_accounts"]["accounts"]

    assert len(accounts) == 2
    assert all(account["account_id_basis"] == "ACCOUNT_ALIAS_MISSING_AMBIGUOUS_MULTI_ROW" for account in accounts)
    assert len({account["account_id"] for account in accounts}) == 2
    codes = {warning["code"] for warning in result["ledger"]["reconciliation_warnings"]}
    assert "ACCOUNT_ALIAS_MISSING_FOR_MULTI_ACCOUNT_ROW" in codes


def test_no_account_snapshot_sheet_yields_no_accounts_and_not_provided_aggregate(tmp_path: Path):
    result = import_workbook(workbook_path=_workbook(tmp_path / "no-accounts.xlsx"), portfolio_root=tmp_path / "private")

    assert result["investment_accounts"]["accounts"] == []
    assert result["investment_accounts_portfolio_context"]["status"] == "NOT_PROVIDED"
    assert result["investment_accounts_portfolio_context"]["account_count"] == 0


def test_investment_account_import_is_deterministic_and_content_addressed(tmp_path: Path):
    workbook_path = _multi_account_workbook(tmp_path / "multi.xlsx")
    first = import_workbook(workbook_path=workbook_path, portfolio_root=tmp_path / "private")
    second = import_workbook(workbook_path=workbook_path, portfolio_root=tmp_path / "private")

    assert second["status"] == "REUSED_IDENTICAL"
    assert first["investment_accounts"]["artifact_identity"] == second["investment_accounts"]["artifact_identity"]
    assert first["investment_accounts_portfolio_context"]["artifact_identity"] == second["investment_accounts_portfolio_context"]["artifact_identity"]
